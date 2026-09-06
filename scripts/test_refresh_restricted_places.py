import json
import math
import os
import sys
import tempfile
import unittest
from collections import Counter
from datetime import datetime
from pathlib import Path
from unittest import mock

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import refresh_restricted_places as builder  # noqa: E402


OFFICIAL_WORKBOOK_VALUE = os.environ.get("NJ_ABC_TEST_WORKBOOK")
OFFICIAL_WORKBOOK = Path(OFFICIAL_WORKBOOK_VALUE) if OFFICIAL_WORKBOOK_VALUE else None


class EssexLookup:
    def find(self, lat: float, lon: float) -> str | None:
        return "Essex" if builder.valid_coordinate(lat, lon) else None


class ABCImportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workbook_path = Path(self.temp_dir.name) / "abc-retail.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Retail License Listing", None, None, None, None, None, None, None])
        sheet.append(["07/01/2026", None, None, None, None, None, None, None])
        sheet.append([
            "License Number", "License Type", "Establishment", "Licensee",
            "Inactivity Start Date", "Effective Date", "City", "Premise Address",
        ])
        sheet.append([
            "0707-33-004-007", "Plenary Retail Consumption License", "DANAHERS",
            "BRIBEN LLC", None, datetime(2026, 1, 20), "FAIRFIELD", "31 PASSAIC AVE",
        ])
        sheet.append([
            "0704-33-007-007", "Plenary Retail Consumption License",
            "GRASSHOPPER BAR & RESTAURANT", "EDWARD FITZPATRICK INC", None,
            datetime(2025, 7, 1), "CEDAR GROVE", "292 GROVE AVENUE",
        ])
        sheet.append([
            "0707-33-999-001", "Plenary Retail Consumption License", "INACTIVE BAR",
            "INACTIVE LLC", datetime(2025, 8, 1), datetime(2025, 7, 1),
            "FAIRFIELD", "1 CLOSED ROAD",
        ])
        sheet.append([
            "0707-44-999-001", "Plenary Retail Distribution License", "PACKAGE STORE",
            "PACKAGE LLC", None, datetime(2025, 7, 1), "FAIRFIELD", "2 PACKAGE ROAD",
        ])
        sheet.append([
            "0707-33-997-001", "Plenary Retail Consumption License", "CITY IN ADDRESS",
            "ADDRESS LLC", None, datetime(2025, 7, 1), "",
            "10 MAIN STREET FAIRFIELD, NJ 07004",
        ])
        sheet.append([
            "0707-33-998-001", "Plenary Retail Consumption License", "MISSING ADDRESS",
            "MISSING LLC", None, datetime(2025, 7, 1), "FAIRFIELD", "",
        ])
        sheet.append([
            "0707-35-996-001", "Seasonal Retail Cons Lic 11/15-4/30", "WINTER CLUB",
            "WINTER CLUB LLC", None, datetime(2025, 7, 1), "FAIRFIELD", "3 WINTER ROAD",
        ])
        sheet.append([
            "0707-34-994-001", "Seasonal Retail Cons Lic 7/1-11/14 and 5/1-6/30",
            "SUMMER CLUB", "SUMMER CLUB LLC", None, datetime(2025, 7, 1),
            "FAIRFIELD", "5 SUMMER ROAD",
        ])
        sheet.append([
            "0707-33-995-001", "Plenary Retail Consumption License", "Not Available",
            "FRIENDLY LICENSEE LLC", None, datetime(2025, 7, 1), "FAIRFIELD", "4 MAIN ROAD",
        ])
        workbook.save(self.workbook_path)
        workbook.close()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_parser_keeps_active_consumption_and_excludes_inactive_distribution(self) -> None:
        audit: Counter[str] = Counter()
        rows = builder.read_abc_retail_rows(self.workbook_path, audit=audit)
        self.assertEqual(
            {row["license number"] for row in rows},
            {
                "0707-33-004-007", "0704-33-007-007",
                "0707-33-997-001", "0707-33-998-001",
                "0707-35-996-001", "0707-33-995-001",
                "0707-34-994-001",
            },
        )
        self.assertEqual(
            next(row for row in rows if row["license number"] == "0707-33-997-001")["city"],
            "",
        )
        self.assertEqual(audit["eligibleActiveOnPremises"], 7)
        self.assertEqual(audit["excludedInactiveOnPremises"], 1)
        self.assertEqual(audit["excludedNonOnPremises"], 1)
        self.assertEqual(audit["missingCityField"], 1)
        self.assertEqual(audit["missingPremiseAddress"], 1)

    def test_parser_requires_exactly_one_matching_sheet(self) -> None:
        multiple_path = Path(self.temp_dir.name) / "multiple.xlsx"
        workbook = openpyxl.load_workbook(self.workbook_path)
        duplicate = workbook.create_sheet("Unexpected second retail sheet")
        duplicate.append([
            "License Number", "License Type", "Establishment", "Licensee",
            "Inactivity Start Date", "Effective Date", "City", "Premise Address",
        ])
        duplicate.sheet_state = "hidden"
        workbook.save(multiple_path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "exactly one retail-license header"):
            builder.read_abc_retail_rows(multiple_path)

        repeated_path = Path(self.temp_dir.name) / "repeated-header.xlsx"
        workbook = openpyxl.load_workbook(self.workbook_path)
        workbook.active.append([
            "License Number", "License Type", "Establishment", "Licensee",
            "Inactivity Start Date", "Effective Date", "City", "Premise Address",
        ])
        workbook.save(repeated_path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "found 2 matching headers"):
            builder.read_abc_retail_rows(repeated_path)

        duplicate_column_path = Path(self.temp_dir.name) / "duplicate-column.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append([
            "License Number", "License Type", "Establishment", "Licensee",
            "Inactivity Start Date", "Effective Date", "City", "City",
            "Premise Address",
        ])
        workbook.save(duplicate_column_path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "each required column exactly once"):
            builder.read_abc_retail_rows(duplicate_column_path)

        missing_path = Path(self.temp_dir.name) / "missing.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["not", "a", "retail", "license", "sheet"])
        workbook.save(missing_path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "found 0 matching headers"):
            builder.read_abc_retail_rows(missing_path)

        decoy_path = Path(self.temp_dir.name) / "decoy.xlsx"
        workbook = openpyxl.load_workbook(self.workbook_path)
        decoy = workbook.create_sheet("Incomplete decoy")
        decoy.append(["License Number", "License Type", "City"])
        workbook.save(decoy_path)
        workbook.close()
        self.assertEqual(len(builder.read_abc_retail_rows(decoy_path)), 7)

        for sheet_state in ("hidden", "veryHidden"):
            with self.subTest(sheet_state=sheet_state):
                hidden_path = Path(self.temp_dir.name) / f"{sheet_state}.xlsx"
                workbook = openpyxl.load_workbook(self.workbook_path)
                source = workbook.active
                decoy = workbook.create_sheet("Visible instructions")
                decoy.append(["NJ ABC retail license instructions"])
                source.sheet_state = sheet_state
                workbook.active = workbook.index(decoy)
                workbook.save(hidden_path)
                workbook.close()
                with self.assertRaisesRegex(ValueError, "source sheet must be visible"):
                    builder.read_abc_retail_rows(hidden_path)

    def test_display_name_normalizes_numeric_ordinals(self) -> None:
        self.assertEqual(
            builder.display_name("1ST, 2ND, 3RD & 4TH FLOOR LOUNGE"),
            "1st, 2nd, 3rd & 4th Floor Lounge",
        )
        self.assertEqual(
            builder.display_name("11TH, 12TH, 13TH & 21ST AVENUE BAR"),
            "11th, 12th, 13th & 21st Avenue Bar",
        )
        self.assertEqual(builder.display_name("43RD STREET SOCIAL CLUB"), "43rd Street Social Club")
        self.assertEqual(builder.display_name("ST JOSEPH'S VFW"), "St Joseph's VFW")
        self.assertEqual(builder.display_name("21ST Amendment"), "21ST Amendment")

    def test_materialized_places_are_high_confidence_and_deterministic(self) -> None:
        rows = builder.read_abc_retail_rows(self.workbook_path)
        coordinates = {
            "0707-33-004-007": (40.867052207504, -74.282010831353, "31 PASSAIC AVE, FAIRFIELD, NJ, 07004"),
            "0704-33-007-007": (40.853533672427, -74.231348034666, "292 GROVE AVE, CEDAR GROVE, NJ, 07009"),
            "0707-33-997-001": (40.86, -74.28, "10 MAIN ST, FAIRFIELD, NJ, 07004"),
            "0707-35-996-001": (40.861, -74.281, "3 WINTER RD, FAIRFIELD, NJ, 07004"),
            "0707-34-994-001": (40.863, -74.283, "5 SUMMER RD, FAIRFIELD, NJ, 07004"),
            "0707-33-995-001": (40.862, -74.282, "4 MAIN RD, FAIRFIELD, NJ, 07004"),
        }
        matches = {
            builder.abc_geocode_key(row): builder.geocode_match(
                *coordinates[row["license number"]],
                provider="Test Geocoder",
                matchType="PointAddress",
            )
            for row in rows
            if row["license number"] in coordinates
        }
        first = builder.abc_places_from_rows(rows, matches, EssexLookup())
        second = builder.abc_places_from_rows(rows, matches, EssexLookup())
        shifted_matches = {
            key: builder.geocode_match(
                value["latitude"] + 0.0001,
                value["longitude"] + 0.0001,
                value["address"],
                provider=value["provider"],
                matchType=value["matchType"],
            )
            for key, value in matches.items()
        }
        shifted = builder.abc_places_from_rows(rows, shifted_matches, EssexLookup())

        self.assertEqual([place["id"] for place in first], [place["id"] for place in second])
        self.assertEqual([place["id"] for place in first], [place["id"] for place in shifted])
        self.assertEqual(
            {place["name"] for place in first},
            {
                "Danahers", "Grasshopper Bar & Restaurant", "City In Address",
                "Winter Club", "Friendly Licensee LLC",
                "Summer Club",
            },
        )
        self.assertTrue(all(place["type"] == "Alcohol-Serving Location" for place in first))
        self.assertTrue(all(place["source"] == builder.ABC_SOURCE for place in first))
        self.assertTrue(all(place["confidence"] == "high" for place in first))
        self.assertTrue(all(place["_priority"] == 120 for place in first))
        winter = next(place for place in first if place["name"] == "Winter Club")
        summer = next(place for place in first if place["name"] == "Summer Club")
        self.assertIn("November 15 through April 30", winter["notes"])
        self.assertIn("May 1 through November 14", summer["notes"])
        self.assertIsNone(
            builder.abc_license_code("malformed", "Seasonal Retail Consumption License")
        )

    def test_dedupe_prefers_abc_when_osm_address_matches(self) -> None:
        abc = builder.make_place(
            name="DANAHERS",
            place_type="Alcohol-Serving Location",
            latitude=40.8670522,
            longitude=-74.2820108,
            county="Essex",
            source=builder.ABC_SOURCE,
            source_id="0707-33-004-007",
            address="31 PASSAIC AVE, FAIRFIELD, NJ, 07004",
            priority=120,
        )
        osm = builder.make_place(
            name="Danaher's Pub",
            place_type="Alcohol-Serving Location",
            latitude=40.86745,
            longitude=-74.28201,
            county="Essex",
            source="OpenStreetMap contributors",
            source_id="node/123",
            address="31 Passaic Avenue, Fairfield, NJ, 07004",
            priority=40,
        )

        result = builder.dedupe([osm, abc])
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0]["source"], builder.ABC_SOURCE)

    def test_dedupe_preserves_distinct_abc_licenses_at_one_address(self) -> None:
        places = [
            builder.make_place(
                name=name,
                place_type="Alcohol-Serving Location",
                latitude=40.919,
                longitude=-74.075,
                county="Bergen",
                source=builder.ABC_SOURCE,
                source_id=license_number,
                address="1 GARDEN STATE PLAZA, PARAMUS, NJ, 07652",
                priority=120,
                stable_source_id=True,
            )
            for name, license_number in (
                ("CAPITAL GRILLE", "0246-33-003-013"),
                ("FOGO DE CHAO", "0246-33-009-009"),
            )
        ]

        result = builder.dedupe(places)
        self.assertEqual(len(result), 2)
        self.assertEqual({place["name"] for place in result}, {"CAPITAL GRILLE", "FOGO DE CHAO"})

    def test_dedupe_grid_covers_full_500_meter_name_threshold(self) -> None:
        places = [
            builder.make_place(
                name="Boundary Test School",
                place_type="School",
                latitude=latitude,
                longitude=-74.0,
                county="Essex",
                source="test",
                source_id=str(index),
            )
            for index, latitude in enumerate((40.0019, 40.0049))
        ]
        self.assertLess(builder.distance_meters(*places), 500)
        self.assertTrue(builder.duplicate(*places))
        self.assertEqual(len(builder.dedupe(places)), 1)

        latitude = 41.39
        longitude = -74.2001

        def place_at_distance(distance: float, source_id: str) -> dict:
            longitude_delta = distance / (
                111_320 * math.cos(math.radians(latitude))
            )
            return builder.make_place(
                name="Three Cell Boundary School",
                place_type="School",
                latitude=latitude,
                longitude=longitude - longitude_delta,
                county="Bergen",
                source="test",
                source_id=source_id,
            )

        origin = place_at_distance(0, "origin")
        inside = place_at_distance(499, "inside")
        outside = place_at_distance(501, "outside")
        inside_distance = builder.distance_meters(origin, inside)
        outside_distance = builder.distance_meters(origin, outside)
        self.assertLess(inside_distance, 500)
        self.assertGreater(outside_distance, 500)
        self.assertTrue(builder.duplicate(origin, inside))
        self.assertFalse(builder.duplicate(origin, outside))
        self.assertEqual(
            abs(
                int(origin["longitude"] / builder.DEDUPE_GRID_SIZE_DEGREES)
                - int(inside["longitude"] / builder.DEDUPE_GRID_SIZE_DEGREES)
            ),
            3,
        )
        self.assertEqual(len(builder.dedupe([origin, inside])), 1)
        self.assertEqual(len(builder.dedupe([origin, outside])), 2)

    def test_refresh_paths_cannot_overwrite_pinned_release_artifacts(self) -> None:
        custom_cache = Path(self.temp_dir.name) / "refresh-cache.json"
        custom_report = Path(self.temp_dir.name) / "refresh-unmatched.json"
        cache, report, review = builder.resolve_abc_artifact_paths(
            custom_cache,
            custom_report,
            refresh_geocodes=True,
        )
        self.assertEqual((cache, report, review), (custom_cache, custom_report, None))
        with self.assertRaisesRegex(ValueError, "requires explicit"):
            builder.resolve_abc_artifact_paths(None, custom_report, refresh_geocodes=True)
        with self.assertRaisesRegex(ValueError, "pinned release geocode cache"):
            builder.resolve_abc_artifact_paths(
                builder.ABC_REVIEWED_GEOCODE_CACHE,
                custom_report,
                refresh_geocodes=True,
            )
        with self.assertRaisesRegex(ValueError, "published unmatched report"):
            builder.resolve_abc_artifact_paths(
                custom_cache,
                builder.ABC_UNMATCHED_REPORT,
                refresh_geocodes=True,
            )
        with self.assertRaisesRegex(ValueError, "published unmatched report"):
            builder.resolve_abc_artifact_paths(
                builder.ABC_UNMATCHED_REPORT,
                custom_report,
                refresh_geocodes=True,
            )
        with self.assertRaisesRegex(ValueError, "pinned geocode review"):
            builder.resolve_abc_artifact_paths(
                custom_cache,
                builder.ABC_GEOCODE_REVIEW,
                refresh_geocodes=True,
            )
        with self.assertRaisesRegex(ValueError, "must be different paths"):
            builder.resolve_abc_artifact_paths(
                custom_cache,
                custom_cache,
                refresh_geocodes=True,
            )
        for tracked_output in (
            builder.ROOT / "data" / "v2" / "essex.json",
            builder.ROOT / "nj_county_manifest_v2.json",
            builder.ROOT / "nj_dataset_summary_v2.json",
        ):
            with self.subTest(tracked_output=tracked_output):
                with self.assertRaisesRegex(ValueError, "outside the repository"):
                    builder.resolve_abc_artifact_paths(
                        custom_cache,
                        tracked_output,
                        refresh_geocodes=True,
                    )
        with self.assertRaisesRegex(ValueError, "source workbook"):
            builder.resolve_abc_artifact_paths(
                self.workbook_path,
                custom_report,
                refresh_geocodes=True,
                source_workbook_path=self.workbook_path,
            )
        symlink_cache = Path(self.temp_dir.name) / "release-cache-link.json"
        symlink_cache.symlink_to(builder.ABC_REVIEWED_GEOCODE_CACHE)
        with self.assertRaisesRegex(ValueError, "pinned release geocode cache"):
            builder.resolve_abc_artifact_paths(
                symlink_cache,
                custom_report,
                refresh_geocodes=True,
            )
        self.assertEqual(
            builder.resolve_abc_artifact_paths(None, None, refresh_geocodes=False),
            (
                builder.ABC_REVIEWED_GEOCODE_CACHE,
                builder.ABC_UNMATCHED_REPORT,
                builder.ABC_GEOCODE_REVIEW,
            ),
        )

    def test_refresh_main_changes_only_explicit_staging_pair(self) -> None:
        release_paths = sorted((builder.ROOT / "data" / "v2").glob("*.json")) + [
            builder.ROOT / "nj_county_manifest_v2.json",
            builder.ROOT / "nj_dataset_summary_v2.json",
            builder.ABC_UNMATCHED_REPORT,
            builder.ABC_REVIEWED_GEOCODE_CACHE,
            builder.ABC_GEOCODE_REVIEW,
        ]
        release_hashes = {path: builder.file_sha256(path) for path in release_paths}
        custom_cache = Path(self.temp_dir.name) / "refresh-cache.json"
        custom_report = Path(self.temp_dir.name) / "refresh-unmatched.json"

        def fake_refresh(*args, **kwargs):
            self.assertTrue(kwargs["refresh_geocodes"])
            self.assertFalse(kwargs["require_reviewed_cache"])
            self.assertIsNone(kwargs["geocode_review_path"])
            self.assertEqual(kwargs["geocode_cache_path"], custom_cache)
            self.assertEqual(kwargs["unmatched_report_path"], custom_report)
            custom_cache.write_text('{"matches": {}}\n')
            custom_report.write_text('{"unresolvedLicenses": []}\n')
            return []

        argv = [
            "refresh_restricted_places.py",
            "--abc-retail-file", str(self.workbook_path),
            "--refresh-abc-geocodes",
            "--abc-geocode-cache", str(custom_cache),
            "--abc-unmatched-report", str(custom_report),
        ]
        with (
            mock.patch.object(sys, "argv", argv),
            mock.patch.object(
                builder,
                "ABC_REPORT_SHA256",
                builder.file_sha256(self.workbook_path),
            ),
            mock.patch.object(builder, "fetch_json", return_value={"features": []}),
            mock.patch.object(builder, "load_abc_retail", side_effect=fake_refresh),
        ):
            self.assertEqual(builder.main(), 0)

        self.assertTrue(custom_cache.is_file())
        self.assertTrue(custom_report.is_file())
        self.assertEqual(
            {path: builder.file_sha256(path) for path in release_paths},
            release_hashes,
        )

    def test_refresh_rejects_skip_abc_before_fetching(self) -> None:
        argv = [
            "refresh_restricted_places.py",
            "--refresh-abc-geocodes",
            "--skip-abc",
            "--abc-geocode-cache", str(Path(self.temp_dir.name) / "cache.json"),
            "--abc-unmatched-report", str(Path(self.temp_dir.name) / "report.json"),
        ]
        with (
            mock.patch.object(sys, "argv", argv),
            mock.patch.object(builder, "fetch_json") as fetch_json,
            self.assertRaises(SystemExit),
        ):
            builder.main()
        fetch_json.assert_not_called()

    def test_unmatched_report_preserves_missing_and_ungeocoded_rows(self) -> None:
        audit: Counter[str] = Counter()
        rows = builder.read_abc_retail_rows(self.workbook_path, audit=audit)
        danahers = next(row for row in rows if row["license number"] == "0707-33-004-007")
        matches = {
            builder.abc_geocode_key(danahers): builder.geocode_match(
                40.867052207504,
                -74.282010831353,
                "31 PASSAIC AVE, FAIRFIELD, NJ, 07004",
                provider="US Census Bureau",
                matchType="Exact",
            ),
        }
        report_path = Path(self.temp_dir.name) / "unmatched.json"
        builder.write_abc_unmatched_report(
            report_path,
            self.workbook_path,
            rows,
            matches,
            audit,
        )

        report = json.loads(report_path.read_text())
        self.assertNotIn("geocodeReview", report)
        self.assertNotIn("geocodeReviewSHA256", report)
        reasons = {
            row["licenseNumber"]: row["reason"]
            for row in report["unresolvedLicenses"]
        }
        self.assertEqual(reasons["0707-33-998-001"], "missingPremiseAddress")
        self.assertEqual(reasons["0704-33-007-007"], "noAcceptedGeocode")
        self.assertEqual(report["audit"]["geocoded"], 1)
        self.assertEqual(report["audit"]["geocodedByProvider"], {"US Census Bureau": 1})

    def test_county_prefix_validation_rejects_wrong_pin(self) -> None:
        rows = builder.read_abc_retail_rows(self.workbook_path)
        danahers = next(row for row in rows if row["license number"] == "0707-33-004-007")
        key = builder.abc_geocode_key(danahers)

        class WrongCountyLookup:
            def find(self, lat: float, lon: float) -> str | None:
                return "Hunterdon"

        accepted, rejected = builder.validate_abc_geocode_matches(
            [danahers],
            {key: builder.geocode_match(40.63, -74.91, provider="US Census Bureau")},
            WrongCountyLookup(),
        )
        self.assertEqual(accepted, {})
        self.assertEqual(rejected[key][0], "coordinateCountyMismatch")

    def test_provider_quality_gate_rejects_nonexact_and_low_score_matches(self) -> None:
        rows = builder.read_abc_retail_rows(self.workbook_path)
        danahers = next(row for row in rows if row["license number"] == "0707-33-004-007")
        key = builder.abc_geocode_key(danahers)
        base = (40.8670522, -74.2820108, "31 PASSAIC AVE, FAIRFIELD, NJ")

        candidates = {
            "census_nonexact": builder.geocode_match(
                *base, provider="US Census Bureau", matchType="Non_Exact"
            ),
            "njgin_low_score": builder.geocode_match(
                *base, provider="NJ Office of GIS NJ_Geocode", status="M", score=94,
                matchType="PointAddress", region="NJ", county="Essex",
            ),
            "njgin_accepted": builder.geocode_match(
                *base, provider="NJ Office of GIS NJ_Geocode", status="M", score=98,
                matchType="PointAddress", region="NJ", county="Essex",
            ),
        }
        expected = {
            "census_nonexact": "censusMatchNotExact",
            "njgin_low_score": "njginScoreBelow95",
        }
        for label, reason in expected.items():
            with self.subTest(label=label):
                accepted, rejected = builder.validate_abc_geocode_matches(
                    [danahers], {key: candidates[label]}, EssexLookup()
                )
                self.assertEqual(accepted, {})
                self.assertEqual(rejected[key][0], reason)
        accepted, rejected = builder.validate_abc_geocode_matches(
            [danahers], {key: candidates["njgin_accepted"]}, EssexLookup()
        )
        self.assertIn(key, accepted)
        self.assertEqual(rejected, {})

    def test_address_number_and_provider_county_normalization(self) -> None:
        self.assertEqual(builder.leading_address_numbers("160 & 170 MAIN STREET"), {160, 170})
        self.assertEqual(builder.leading_address_numbers("121 & 131 BROADWAY"), {121, 131})
        self.assertEqual(builder.leading_address_numbers("196 198 MARKET STREET"), {196, 198})
        self.assertEqual(builder.leading_address_numbers("1312 AND 1318 WEST BRIGANTINE AVENUE"), {1312, 1318})
        self.assertEqual(builder.leading_address_numbers("20,22 AND 24 MAIN STREET"), {20, 22, 24})
        self.assertEqual(
            builder.leading_address_numbers("2002 - 2004 - 2006 - 2008 BOARDWALK"),
            {2002, 2004, 2006, 2008},
        )
        self.assertEqual(
            builder.leading_address_numbers("97-99- 101 MONMOUTH ST"),
            {97, 99, 101},
        )
        self.assertEqual(builder.leading_address_numbers("50-54 WATCHUNG AVENUE"), set(range(50, 55)))
        chained_premises = {
            "229-231 -233 S MAIN STREET": {229, 231, 233},
            "2002 - 2004 - 2006 - 2008 BOARDWALK": {2002, 2004, 2006, 2008},
            "201 - 205 - 209 OLDE NEW JERSEY AVENUE": {201, 205, 209},
            "32-34- 36 BRANFORD PL 211 HALSEY": {32, 34, 36},
            "106-108-110 SCHANCK ROAD": {106, 108, 110},
            "15-17- 19-21 BROAD STREET": {15, 17, 19, 21},
            "97-99- 101 MONMOUTH ST": {97, 99, 101},
            "4-6-8 CHURCH STREET": {4, 6, 8},
        }
        for address, expected in chained_premises.items():
            with self.subTest(address=address):
                self.assertEqual(builder.leading_address_numbers(address), expected)
        self.assertEqual(builder.county_name("Atlantic ..."), "Atlantic")
        self.assertEqual(builder.county_name("ESSEX COUNTY"), "Essex")
        self.assertEqual(
            builder.abc_geocoder_components(
                {
                    "premise address": "341 SOUTH AVENUE EAST WESTFIELD, NJ 07090",
                    "city": "",
                },
                {"WESTFIELD", "ATLANTIC CITY"},
            ),
            ("341 SOUTH AVENUE EAST", "WESTFIELD", "NJ", "07090"),
        )
        self.assertEqual(
            builder.abc_geocoder_components(
                {
                    "premise address": "10 MAIN STREET FAIRFIELD, NEW JERSEY 07004",
                    "city": "FAIRFIELD",
                },
                {"FAIRFIELD"},
            ),
            ("10 MAIN STREET", "FAIRFIELD", "NJ", "07004"),
        )

    def test_census_response_completeness_validation(self) -> None:
        builder.validate_census_response(
            [
                ["a", "input", "Match", "Exact", "1 MAIN ST", "-74.2,40.8"],
                ["b", "input", "No_Match"],
            ],
            ["a", "b"],
        )
        with self.assertRaises(ValueError):
            builder.validate_census_response([["a", "input", "Match"]], ["a"])
        with self.assertRaises(ValueError):
            builder.validate_census_response(
                [["a", "input", "Match", "Exact", "1 MAIN ST", "not-coordinates"]],
                ["a"],
            )
        with self.assertRaises(ValueError):
            builder.validate_census_response(
                [["a", "input", "No_Match"], ["a", "input", "No_Match"]],
                ["a", "b"],
            )

    def test_geocode_cache_round_trip_preserves_provider_review_metadata(self) -> None:
        cache_path = Path(self.temp_dir.name) / "geocodes.json"
        matches = {
            "license|address|city": builder.geocode_match(
                40.8,
                -74.2,
                "1 MAIN ST, TEST, NJ",
                provider="NJ Office of GIS NJ_Geocode",
                score=100,
                matchType="PointAddress",
                county="Essex",
            )
        }
        builder.save_geocode_cache(cache_path, matches)
        loaded = builder.load_geocode_cache(cache_path)
        self.assertEqual(loaded, matches)
        payload = json.loads(cache_path.read_text())
        self.assertEqual(payload["version"], 2)
        self.assertEqual(payload["providers"], ["NJ Office of GIS NJ_Geocode"])

    def test_abc_only_augmentation_preserves_non_alcohol_baseline(self) -> None:
        school = builder.make_place(
            name="Test School", place_type="School", latitude=40.8, longitude=-74.2,
            county="Essex", source="baseline", source_id="school", priority=90,
        )
        osm = builder.make_place(
            name="Danaher's Pub", place_type="Alcohol-Serving Location",
            latitude=40.86745, longitude=-74.28201, county="Essex",
            source="OpenStreetMap contributors", source_id="node/123",
            address="31 Passaic Avenue, Fairfield, NJ, 07004", priority=40,
        )
        abc = builder.make_place(
            name="Danahers", place_type="Alcohol-Serving Location",
            latitude=40.8670522, longitude=-74.2820108, county="Essex",
            source=builder.ABC_SOURCE, source_id="0707-33-004-007",
            address="31 PASSAIC AVE, FAIRFIELD, NJ, 07004", priority=120,
            stable_source_id=True,
        )
        result, replaced = builder.augment_v2_baseline_with_abc([school, osm], [abc])
        self.assertEqual(replaced, 1)
        self.assertIn(school, result)
        self.assertIn(abc, result)
        self.assertNotIn(osm, result)


@unittest.skipUnless(
    OFFICIAL_WORKBOOK is not None and OFFICIAL_WORKBOOK.exists(),
    "set NJ_ABC_TEST_WORKBOOK to run the official-workbook regression test",
)
class OfficialABCWorkbookTests(unittest.TestCase):
    def test_july_2026_counts_and_known_essex_licenses(self) -> None:
        audit: Counter[str] = Counter()
        rows = builder.read_abc_retail_rows(OFFICIAL_WORKBOOK, audit=audit)
        by_license = {row["license number"]: row for row in rows}

        self.assertEqual(len(rows), 5_774)
        self.assertEqual(sum(bool(row["premise address"]) for row in rows), 5_770)
        self.assertEqual(audit["excludedInactiveOnPremises"], 1_189)
        self.assertEqual(audit["missingCityField"], 55)
        self.assertEqual(audit["missingPremiseAddress"], 4)
        self.assertTrue(all(not builder.clean(row["inactivity start date"]) for row in rows))
        self.assertTrue(all(row["license code"] not in {"43", "44"} for row in rows))

        grasshopper = by_license["0704-33-007-007"]
        self.assertEqual(grasshopper["establishment"], "GRASSHOPPER BAR & RESTAURANT")
        self.assertEqual(grasshopper["city"], "CEDAR GROVE")
        self.assertEqual(grasshopper["premise address"], "292 GROVE AVENUE")

        danahers = by_license["0707-33-004-007"]
        self.assertEqual(danahers["establishment"], "DANAHERS")
        self.assertEqual(danahers["city"], "FAIRFIELD")
        self.assertEqual(danahers["premise address"], "31 PASSAIC AVE")


class GeneratedReleaseArtifactTests(unittest.TestCase):
    def test_v3_outputs_reconcile_and_include_required_essex_markers(self) -> None:
        places: list[dict] = []
        for county in builder.COUNTIES:
            path = builder.ROOT / "data" / "v2" / f"{county.lower().replace(' ', '_')}.json"
            county_places = json.loads(path.read_text())
            self.assertTrue(all(place["county"] == county for place in county_places))
            places.extend(county_places)

        ids = [place["id"] for place in places]
        self.assertEqual(len(ids), len(set(ids)))
        abc_places = [place for place in places if place.get("source") == builder.ABC_SOURCE]
        self.assertEqual(len(abc_places), builder.ABC_REVIEWED_GEOCODE_COUNT)

        for license_number, (expected_name, expected_county) in builder.ABC_REGRESSION_LICENSES.items():
            expected_id = builder.deterministic_source_id(
                builder.ABC_SOURCE, license_number, "Alcohol-Serving Location"
            )
            matches = [place for place in places if place["id"] == expected_id]
            self.assertEqual(len(matches), 1)
            self.assertEqual(matches[0]["name"], expected_name)
            self.assertEqual(matches[0]["county"], expected_county)
            self.assertEqual(matches[0]["type"], "Alcohol-Serving Location")

        summary = json.loads((builder.ROOT / "nj_dataset_summary_v2.json").read_text())
        self.assertEqual(summary["version"], builder.DATASET_REVISION)
        self.assertEqual(summary["totalPlaces"], len(places))
        self.assertEqual(
            summary["countsByType"],
            dict(sorted(Counter(place["type"] for place in places).items())),
        )
        self.assertEqual(summary["countsBySource"][builder.ABC_SOURCE], len(abc_places))

        report = json.loads((builder.ROOT / "nj_abc_unmatched_v2.json").read_text())
        self.assertEqual(report["audit"]["materialized"], len(abc_places))
        self.assertEqual(report["audit"]["geocoded"], len(abc_places))
        self.assertEqual(report["audit"]["unresolved"], len(report["unresolvedLicenses"]))
        self.assertEqual(report["geocodeReview"], builder.ABC_GEOCODE_REVIEW.name)
        self.assertEqual(report["geocodeReviewSHA256"], builder.ABC_GEOCODE_REVIEW_SHA256)

        provenance = summary["sourceProvenance"]["njABC"]
        self.assertEqual(provenance["geocodeReview"], builder.ABC_GEOCODE_REVIEW.name)
        self.assertEqual(provenance["geocodeReviewSHA256"], builder.ABC_GEOCODE_REVIEW_SHA256)

        manifest = json.loads((builder.ROOT / "nj_county_manifest_v2.json").read_text())
        self.assertEqual(manifest["version"], builder.DATASET_REVISION)
        self.assertEqual(len(manifest["counties"]), 21)
        self.assertTrue(all(
            entry["version"] == builder.DATASET_REVISION
            and entry["datasetURL"].endswith(f"?v={builder.DATASET_REVISION}")
            for entry in manifest["counties"]
        ))

        self.assertEqual(
            builder.file_sha256(builder.ABC_REVIEWED_GEOCODE_CACHE),
            builder.ABC_REVIEWED_GEOCODE_CACHE_SHA256,
        )
        self.assertEqual(
            builder.file_sha256(builder.ABC_GEOCODE_REVIEW),
            builder.ABC_GEOCODE_REVIEW_SHA256,
        )


if __name__ == "__main__":
    unittest.main()
